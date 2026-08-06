#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""手術式為 my錢錢.xlsx 新增『投資報酬』分頁：分期(自訂區間)計算投資報酬，排除中間買賣的影響。
每期只手填兩格：期末日期、期末證券市值(台幣)。期初自動接上一期，買/賣/股息自動從交易明細抓。
  投資損益 = 期末市值 − 期初市值 − 淨投入(買−賣) + 股息      ← 排除加減碼後真正賺的錢
  期間報酬率 = 投資損益 / 平均資本(期初 + 淨投入×0.5)         ← Modified Dietz(期間中點假設)
  年化報酬率 = (1+期間報酬率)^(365/天數) − 1
其餘分頁與資料完全不動。存檔會丟棄『儀表板』圖表 → 之後跑 dashboard_update.py 重建。
用法：python add_investment_return.py
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

LIVE = "/Users/ucpc/Google Drive/我的雲端硬碟/my錢錢.xlsx"
NROWS = 30

TITLE   = Font(name="Arial", size=16, bold=True, color="1F3864")
NOTE    = Font(name="Arial", size=9, italic=True, color="808080")
HEAD_M  = Font(name="Arial", size=10, bold=True, color="7F6000")
HEAD_AW = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FILL_MANHEAD = PatternFill("solid", fgColor="FFE599")
FILL_AUTHEAD = PatternFill("solid", fgColor="808080")
FILL_MAN     = PatternFill("solid", fgColor="FFF9E6")
FILL_AUT     = PatternFill("solid", fgColor="F0F0F0")
FILL_OK      = PatternFill("solid", fgColor="C6EFCE")
FILL_BAD     = PatternFill("solid", fgColor="FFC7CE")
CTR  = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="D9D9D9"); BORDER = Border(THIN, THIN, THIN, THIN)
NT = '#,##0'; PCT = '0.0%'; DAY = '0'; DATE = 'yyyy-mm-dd'

# 交易明細範圍
TA = "'交易明細'!$A$5:$A$515"    # 日期
TL = "'交易明細'!$L$5:$L$515"    # 台幣金額
TC = "'交易明細'!$C$5:$C$515"    # 類型
HELD = "SUM('持股明細'!$I$5:$I$57)"   # 目前證券總市值(台幣)

wb = openpyxl.load_workbook(LIVE)
if "投資報酬" in wb.sheetnames:
    wb.remove(wb["投資報酬"])
pos = wb.sheetnames.index("對帳") + 1 if "對帳" in wb.sheetnames else len(wb.sheetnames)
ws = wb.create_sheet("投資報酬", pos)

ws["A1"] = "📈 投資報酬（分期計算，排除中間買賣的影響）"; ws["A1"].font = TITLE
ws["A2"] = "✍️ 黃底 = 手填     🔒 灰底 = 自動計算（勿改）"; ws["A2"].font = NOTE
ws["A3"] = ("💡 每期只填兩格：①期末日期 ②期末證券市值(台幣，從『持股明細』台幣市值總和貼『數值』凍結)。"
            "期初自動接上一期，買/賣/股息自動從交易明細抓。第一列為期初基準（無損益），第二列起才有報酬。")
ws["A3"].font = NOTE

# 欄位：(標題, man/aut, 寬, 數字格式)
cols = [
    ("✍️ 期末日期", "man", 12, DATE),
    ("✍️ 期末證券市值", "man", 14, NT),
    ("🔒 期初日期", "aut", 12, DATE),
    ("🔒 期初證券市值", "aut", 14, NT),
    ("🔒 本期買入", "aut", 12, NT),
    ("🔒 本期賣出", "aut", 12, NT),
    ("🔒 本期股息", "aut", 11, NT),
    ("🔒 淨投入", "aut", 12, NT),
    ("🔒 投資損益", "aut", 13, NT),
    ("🔒 期間天數", "aut", 9, DAY),
    ("🔒 平均資本", "aut", 13, NT),
    ("🔒 期間報酬率", "aut", 11, PCT),
    ("🔒 年化報酬率", "aut", 11, PCT),
]
# 欄索引
idx = {t.split()[-1]: i for i, (t, *_ ) in enumerate(cols, start=1)}
C_END, C_ENDV, C_BEG, C_BEGV, C_BUY, C_SELL, C_DIV, C_NET, C_PL, C_DAYS, C_CAP, C_R, C_AR = range(1, 14)
Lc = lambda c: get_column_letter(c)

HROW = 4
for i, (title, kind, w, fmt) in enumerate(cols, start=1):
    c = ws.cell(HROW, i, title)
    c.font = HEAD_M if kind == "man" else HEAD_AW
    c.fill = FILL_MANHEAD if kind == "man" else FILL_AUTHEAD
    c.alignment = WRAP; c.border = BORDER
    ws.column_dimensions[Lc(i)].width = w

DATA0 = HROW + 1
for r in range(DATA0, DATA0 + NROWS):
    A, B = f"$A{r}", f"$B{r}"
    if r > DATA0:
        ws.cell(r, C_BEG,  f'=IF($A{r}="","",$A{r-1})')
        ws.cell(r, C_BEGV, f'=IF($A{r}="","",$B{r-1})')
    beg = f"{Lc(C_BEG)}{r}"; begv = f"{Lc(C_BEGV)}{r}"
    # 買/賣/股息：抓 (期初日期, 期末日期] 區間內的交易
    def sif(typ):
        return (f'SUMIFS({TL},{TA},">"&{beg},{TA},"<="&{A},{TC},"{typ}")')
    ws.cell(r, C_BUY,  f'=IF(OR($A{r}="",{beg}=""),"",{sif("買")})')
    ws.cell(r, C_SELL, f'=IF(OR($A{r}="",{beg}=""),"",{sif("賣")})')
    ws.cell(r, C_DIV,  f'=IF(OR($A{r}="",{beg}=""),"",{sif("股息")})')
    buy = f"{Lc(C_BUY)}{r}"; sell = f"{Lc(C_SELL)}{r}"; div = f"{Lc(C_DIV)}{r}"
    net = f"{Lc(C_NET)}{r}"; days = f"{Lc(C_DAYS)}{r}"; cap = f"{Lc(C_CAP)}{r}"
    pl = f"{Lc(C_PL)}{r}"; rate = f"{Lc(C_R)}{r}"
    ws.cell(r, C_NET,  f'=IF(OR($A{r}="",{beg}=""),"",N({buy})-N({sell}))')
    ws.cell(r, C_PL,   f'=IF(OR($A{r}="",{beg}=""),"",{B}-{begv}-{net}+N({div}))')
    ws.cell(r, C_DAYS, f'=IF(OR($A{r}="",{beg}=""),"",{A}-{beg})')
    ws.cell(r, C_CAP,  f'=IF(OR($A{r}="",{beg}=""),"",{begv}+{net}*0.5)')
    ws.cell(r, C_R,    f'=IF(OR($A{r}="",{beg}="",{cap}=0),"",IFERROR({pl}/{cap},""))')
    ws.cell(r, C_AR,   f'=IF(OR({rate}="",{days}="",{days}=0),"",IFERROR((1+{rate})^(365/{days})-1,""))')
    # 樣式
    for i, (title, kind, w, fmt) in enumerate(cols, start=1):
        cell = ws.cell(r, i)
        cell.fill = FILL_MAN if kind == "man" else FILL_AUT
        cell.border = BORDER
        cell.number_format = fmt
        if i in (C_END, C_BEG, C_DAYS): cell.alignment = CTR

# 損益 / 報酬率 綠正紅負
rng_pl = f"{Lc(C_PL)}{DATA0}:{Lc(C_PL)}{DATA0+NROWS-1}"
for col in (C_PL, C_R, C_AR):
    rng = f"{Lc(col)}{DATA0}:{Lc(col)}{DATA0+NROWS-1}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_OK))
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=FILL_BAD))

ws.freeze_panes = "C5"   # 捲動保留 期末日期 + 期末市值

wb.save(LIVE)
print("✅ 已存檔（投資報酬 已建）")
print("分頁順序:", wb.sheetnames)
print("欄位: 期末日期|期末市值|期初日期|期初市值|買|賣|股息|淨投入|投資損益|天數|平均資本|期間報酬率|年化報酬率")
