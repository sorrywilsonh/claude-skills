#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""手術式重建 my錢錢.xlsx 的『儀表板』分頁；其餘分頁與資料完全不動。
版面：KPI(總淨資產/現金總額/股票總額) → 各幣別現金&股票表 → 兩張餅圖 → 手動說明(最下方)。
注意：openpyxl 載入會丟棄舊圖表（圖表僅存在於儀表板，本來就要重建），故安全。
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import PieChart, Reference
from openpyxl.utils import get_column_letter

LIVE = "/Users/ucpc/Google Drive/我的雲端硬碟/my錢錢.xlsx"

# 樣式
TITLE = Font(name="Arial", size=20, bold=True, color="1F3864")
NOTE  = Font(name="Arial", size=9, italic=True, color="808080")
HEADW = Font(name="Arial", size=10, bold=True, color="FFFFFF")
SECT  = Font(name="Arial", size=12, bold=True, color="1F3864")
KPIV  = Font(name="Arial", size=16, bold=True, color="1F3864")
MANF  = Font(name="Arial", size=10, color="7F6000")
BOLD  = Font(name="Arial", size=10, bold=True, color="333333")
FILL_HEAD = PatternFill("solid", fgColor="808080")
FILL_KPI  = PatternFill("solid", fgColor="DDEBF7")
FILL_MAN  = PatternFill("solid", fgColor="FFF9E6")
FILL_LEG  = PatternFill("solid", fgColor="FCE4D6")
CTR = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="D9D9D9"); BORDER = Border(THIN, THIN, THIN, THIN)
NT = '#,##0'; NT2 = '#,##0.00'; PCT = '0.0%'

# 跨表範圍
AB='\'帳戶總覽\'!$B$5:$B$25'; AC='\'帳戶總覽\'!$C$5:$C$25'
AD='\'帳戶總覽\'!$D$5:$D$25'; AF='\'帳戶總覽\'!$F$5:$F$25'
HD='\'持股明細\'!$D$5:$D$57'; HH='\'持股明細\'!$H$5:$H$57'; HI='\'持股明細\'!$I$5:$I$57'

wb = openpyxl.load_workbook(LIVE)
# 移除舊儀表板，於最前面重建
if "儀表板" in wb.sheetnames:
    wb.remove(wb["儀表板"])
d = wb.create_sheet("儀表板", 0)

d["A1"] = "📊 my錢錢 — 資產儀表板"; d["A1"].font = TITLE
d["A2"] = "即時股價/匯率在 Google Sheets 開啟時自動生效。"; d["A2"].font = NOTE
for col, w in [("A",14),("B",13),("C",13),("D",14),("E",13),("F",13),("G",14),("H",13),("J",10),("K",14)]:
    d.column_dimensions[col].width = w

# ---- KPI（總淨資產 / 現金總額 / 股票總額，皆台幣）----
kpis = [("總淨資產 (台幣)", "='帳戶總覽'!$F$27", 1),
        ("現金總額 (台幣)", f"=SUMIFS({AF},{AB},\"銀行\")", 4),
        ("股票總額 (台幣)", f"=SUM({HI})", 7)]
for label, formula, col in kpis:
    lc = d.cell(row=4, column=col, value=label); lc.font = HEADW; lc.fill = FILL_HEAD; lc.alignment = CTR
    d.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col+1)
    vc = d.cell(row=5, column=col, value=formula); vc.font = KPIV; vc.fill = FILL_KPI
    vc.number_format = NT; vc.alignment = CTR
    d.merge_cells(start_row=5, start_column=col, end_row=6, end_column=col+1)

# ---- 各幣別現金 / 股票 ----
d["A8"] = "💵 現金資產（各幣別）"; d["A8"].font = SECT
d["E8"] = "📈 股票資產（各幣別）"; d["E8"].font = SECT
for c, t in [(1,"幣別"),(2,"金額(原幣)"),(3,"(台幣)")]:
    x=d.cell(row=9,column=c,value=t); x.font=HEADW; x.fill=FILL_HEAD; x.alignment=CTR; x.border=BORDER
for c, t in [(5,"幣別"),(6,"金額(原幣)"),(7,"(台幣)")]:
    x=d.cell(row=9,column=c,value=t); x.font=HEADW; x.fill=FILL_HEAD; x.alignment=CTR; x.border=BORDER
for i, cur in enumerate(["TWD","USD","JPY"]):
    r = 10 + i
    # 現金
    d.cell(row=r, column=1, value=cur).alignment = CTR
    d.cell(row=r, column=2, value=f"=SUMIFS({AD},{AB},\"銀行\",{AC},\"{cur}\")").number_format = NT2
    d.cell(row=r, column=3, value=f"=SUMIFS({AF},{AB},\"銀行\",{AC},\"{cur}\")").number_format = NT
    # 股票
    d.cell(row=r, column=5, value=cur).alignment = CTR
    d.cell(row=r, column=6, value=f"=SUMIFS({HH},{HD},\"{cur}\")").number_format = NT2
    d.cell(row=r, column=7, value=f"=SUMIFS({HI},{HD},\"{cur}\")").number_format = NT
    for c in (1,2,3,5,6,7): d.cell(row=r, column=c).border = BORDER

# ---- 餅圖資料：現金 vs 股票（台幣）----
d["J4"] = "類別"; d["K4"] = "台幣金額"
for c in ("J4","K4"): d[c].font = HEADW; d[c].fill = FILL_HEAD; d[c].alignment = CTR
d["J5"] = "現金"; d["K5"] = f"=SUMIFS({AF},{AB},\"銀行\")"; d["K5"].number_format = NT
d["J6"] = "股票"; d["K6"] = f"=SUM({HI})"; d["K6"].number_format = NT

# ---- 圖 1：現金 vs 股票 比例 ----
pie1 = PieChart(); pie1.title = "現金 vs 股票 比例（台幣）"; pie1.height = 8; pie1.width = 10
pie1.add_data(Reference(d, min_col=11, min_row=4, max_row=6), titles_from_data=True)
pie1.set_categories(Reference(d, min_col=10, min_row=5, max_row=6))
pie1.dataLabels = openpyxl.chart.label.DataLabelList(); pie1.dataLabels.showPercent = True
d.add_chart(pie1, "A15")

# ---- 圖 2：全部股票餅圖（依台幣市值）----
pie2 = PieChart(); pie2.title = "全部股票配置（台幣市值）"; pie2.height = 8; pie2.width = 11
pie2.add_data(Reference(wb["持股明細"], min_col=9, min_row=4, max_row=57), titles_from_data=True)
pie2.set_categories(Reference(wb["持股明細"], min_col=2, min_row=5, max_row=57))
pie2.dataLabels = openpyxl.chart.label.DataLabelList(); pie2.dataLabels.showPercent = True
d.add_chart(pie2, "F15")

# ---- 手動更新說明（移到圖形下方）----
d["A33"] = "🔧 需要你『手動更新』的項目（黃底欄位）"; d["A33"].font = BOLD; d["A33"].fill = FILL_LEG
items = [
    "① 設定：建立『子帳戶清單』(帳戶／類型／子分類) — 幣別與識別碼自動，一次性",
    "② 帳戶總覽：選子帳戶 + 填『現金餘額(原幣)』(含各家卡債餘額) — 每月更新當下餘額",
    "③ 持股明細：新增持股填『帳戶／代號／名稱』；基金填『手動現價』",
    "④ 交易明細：每筆證券買/賣/股息記一列 — 投資主力",
    "⑤ 資產快照：每月把上方『總淨資產』貼成數值一列",
    "⑥ 對帳：每月填各帳戶餘額(台幣)＋消費(panel當月總額)＋收入 → 自動比對現金缺漏",
    "※ 日常消費改用 daily-panel 記帳，不在本表登打",
]
for i, txt in enumerate(items, start=34):
    c = d.cell(row=i, column=1, value=txt); c.font = MANF; c.fill = FILL_MAN
    d.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)

wb.save(LIVE)
print("✅ 已更新儀表板:", LIVE)
print("分頁順序:", wb.sheetnames)
