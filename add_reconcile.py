#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""手術式改造 my錢錢.xlsx：
  ① 刪除已無資料來源的『現金流』分頁。
  ② 新增『對帳』分頁：每月比對「現金淨額實際變化」vs「應該的變化」(收入−消費−投入證券+股息)，
     抓出兩方(xlsx帳戶餘額 / panel消費)的缺漏。分帳戶級：每個銀行/卡債各一欄，可定位是哪個帳戶。
其餘分頁與資料完全不動。存檔會丟棄『儀表板』圖表 → 之後跑 dashboard_update.py 重建。
用法：python add_reconcile.py
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

LIVE = "/Users/ucpc/Google Drive/我的雲端硬碟/my錢錢.xlsx"
TOL = 1000          # 差額容忍門檻(台幣)：|差額|<=此值視為吻合
NROWS = 30          # 預先鋪好公式的月份列數

# ---- 樣式（沿用 build_my_money.py 慣例）----
TITLE   = Font(name="Arial", size=16, bold=True, color="1F3864")
NOTE    = Font(name="Arial", size=9, italic=True, color="808080")
HEAD_M  = Font(name="Arial", size=10, bold=True, color="7F6000")   # 手填欄表頭字
HEAD_AW = Font(name="Arial", size=10, bold=True, color="FFFFFF")   # 自動欄表頭字
FILL_MANHEAD = PatternFill("solid", fgColor="FFE599")
FILL_AUTHEAD = PatternFill("solid", fgColor="808080")
FILL_MAN     = PatternFill("solid", fgColor="FFF9E6")
FILL_AUT     = PatternFill("solid", fgColor="F0F0F0")
FILL_OK      = PatternFill("solid", fgColor="C6EFCE")
FILL_BAD     = PatternFill("solid", fgColor="FFC7CE")
CTR = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="D9D9D9"); BORDER = Border(THIN, THIN, THIN, THIN)
NT = '#,##0'

TX = "'交易明細'!$L$5:$L$515"    # 台幣金額
TXM = "'交易明細'!$M$5:$M$515"   # 手續費稅(台幣)
TXN = "'交易明細'!$N$5:$N$515"   # 年月
TXC = "'交易明細'!$C$5:$C$515"   # 類型

wb = openpyxl.load_workbook(LIVE)

# ① 刪 現金流
if "現金流" in wb.sheetnames:
    wb.remove(wb["現金流"])
    print("已刪除『現金流』分頁")

# 由『帳戶總覽』動態取現金/卡債帳戶（名稱含『銀行』=現金；含『卡債』=負債），自動適應日後增減帳戶
ov = wb["帳戶總覽"]
BANKS, DEBTS = [], []
for r in range(5, 26):
    name = ov.cell(r, 1).value
    if not name:
        continue
    if "卡債" in name:
        DEBTS.append(name)
    elif "銀行" in name:
        BANKS.append(name)
print(f"現金帳戶 {len(BANKS)}、卡債帳戶 {len(DEBTS)}")

# ② 建 對帳（放在 資產快照 之後）
if "對帳" in wb.sheetnames:
    wb.remove(wb["對帳"])
pos = wb.sheetnames.index("資產快照") + 1
ws = wb.create_sheet("對帳", pos)

ws["A1"] = "🔍 對帳（每月比對現金實際變化 vs 應該的變化 → 抓兩方缺漏）"; ws["A1"].font = TITLE
ws["A2"] = "✍️ 黃底 = 手填     🔒 灰底 = 自動計算（勿改）"; ws["A2"].font = NOTE
ws["A3"] = ("💡 每月填三種：①各帳戶當月餘額(台幣，可從『帳戶總覽』F欄複製→選擇性貼上數值) "
            "②消費(daily-panel 當月總額) ③收入/薪資。現金淨額＝銀行−卡債；差額≈0 代表兩方吻合。")
ws["A3"].font = NOTE

# 欄位定義：(標題, 種類 man/aut, 寬)
cols = [("✍️ 年月", "man", 9)]
for nm in BANKS:
    cols.append(("✍️ " + nm, "man", 11))
for nm in DEBTS:
    cols.append(("✍️ " + nm, "man", 11))
FIRST_BANK = 2
LAST_BANK = 1 + len(BANKS)
FIRST_DEBT = LAST_BANK + 1
LAST_DEBT = LAST_BANK + len(DEBTS)
auto_cols = ["🔒 現金淨額", "🔒 Δ現金"]
man_after = None
# 之後欄位順序：現金淨額, Δ現金, 投入證券, 股息, 消費, 收入, 預期Δ, 差額, 檢查
tail = [("🔒 現金淨額", "aut", 12), ("🔒 Δ現金", "aut", 11),
        ("🔒 投入證券", "aut", 11), ("🔒 股息", "aut", 10),
        ("✍️ 消費(panel)", "man", 12), ("✍️ 收入/薪資", "man", 12),
        ("🔒 預期Δ", "aut", 11), ("🔒 差額", "aut", 11), ("🔒 檢查", "aut", 13)]
cols += tail

# 欄位索引（1-based）
def col_of(title_sub):
    for i, (t, _, _) in enumerate(cols, start=1):
        if title_sub in t:
            return i
    raise KeyError(title_sub)

C_YM = 1
C_NET = col_of("現金淨額"); C_DELTA = col_of("Δ現金"); C_INV = col_of("投入證券")
C_DIV = col_of("股息"); C_SPEND = col_of("消費"); C_INC = col_of("收入")
C_EXP = col_of("預期Δ"); C_DIFF = col_of("差額"); C_CHK = col_of("檢查")
L = lambda c: get_column_letter(c)

# 表頭（第4列）
HROW = 4
for i, (title, kind, w) in enumerate(cols, start=1):
    c = ws.cell(HROW, i, title)
    c.font = HEAD_M if kind == "man" else HEAD_AW
    c.fill = FILL_MANHEAD if kind == "man" else FILL_AUTHEAD
    c.alignment = WRAP; c.border = BORDER
    ws.column_dimensions[L(i)].width = w

# 資料列公式
DATA0 = HROW + 1
for r in range(DATA0, DATA0 + NROWS):
    net = (f"SUM({L(FIRST_BANK)}{r}:{L(LAST_BANK)}{r})"
           f"-SUM({L(FIRST_DEBT)}{r}:{L(LAST_DEBT)}{r})")
    ws.cell(r, C_NET, f'=IF($A{r}="","",{net})')
    if r > DATA0:
        ws.cell(r, C_DELTA,
                f'=IF(OR($A{r}="",$A{r-1}=""),"",{L(C_NET)}{r}-{L(C_NET)}{r-1})')
    ws.cell(r, C_INV,
            f'=IF($A{r}="","",'
            f'SUMIFS({TX},{TXN},$A{r},{TXC},"買")-SUMIFS({TX},{TXN},$A{r},{TXC},"賣")'
            f'+SUMIFS({TXM},{TXN},$A{r}))')
    ws.cell(r, C_DIV,
            f'=IF($A{r}="","",SUMIFS({TX},{TXN},$A{r},{TXC},"股息"))')
    ws.cell(r, C_EXP,
            f'=IF($A{r}="","",N({L(C_INC)}{r})-N({L(C_SPEND)}{r})-N({L(C_INV)}{r})+N({L(C_DIV)}{r}))')
    ws.cell(r, C_DIFF,
            f'=IF(OR($A{r}="",{L(C_DELTA)}{r}=""),"",{L(C_DELTA)}{r}-{L(C_EXP)}{r})')
    ws.cell(r, C_CHK,
            f'=IF(OR($A{r}="",{L(C_DIFF)}{r}=""),"",'
            f'IF(ABS({L(C_DIFF)}{r})<={TOL},"✅ 吻合","⚠️ 差 "&TEXT({L(C_DIFF)}{r},"#,##0")))')
    # 樣式與數字格式
    for i, (title, kind, w) in enumerate(cols, start=1):
        cell = ws.cell(r, i)
        cell.fill = FILL_MAN if kind == "man" else FILL_AUT
        cell.border = BORDER
        if i != C_YM and i != C_CHK:
            cell.number_format = NT
    ws.cell(r, C_YM).alignment = CTR
    ws.cell(r, C_CHK).alignment = CTR

# 檢查欄條件式上色
chk_range = f"{L(C_CHK)}{DATA0}:{L(C_CHK)}{DATA0+NROWS-1}"
ws.conditional_formatting.add(chk_range,
    FormulaRule(formula=[f'ISNUMBER(SEARCH("⚠",{L(C_CHK)}{DATA0}))'], fill=FILL_BAD))
ws.conditional_formatting.add(chk_range,
    FormulaRule(formula=[f'ISNUMBER(SEARCH("✅",{L(C_CHK)}{DATA0}))'], fill=FILL_OK))

ws.freeze_panes = "B5"   # 捲動時保留『年月』欄

wb.save(LIVE)
print("✅ 已存檔（現金流已刪、對帳已建）")
print("分頁順序:", wb.sheetnames)
print(f"對帳欄位：年月 | {len(BANKS)}現金 | {len(DEBTS)}卡債 | 現金淨額/Δ/投入證券/股息/消費/收入/預期Δ/差額/檢查")
